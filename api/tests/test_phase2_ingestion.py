import pytest
import pandas as pd
import os
import gc
from pathlib import Path
import shutil
from uuid import uuid4
from api.app.services.parsers.tabular import TabularParser
from api.app.services.parsers.knowledge import KnowledgeParser

@pytest.fixture
def temp_storage():
    root = Path("tmp") / f"phase2-ingestion-{uuid4()}"
    root.mkdir(parents=True, exist_ok=True)
    yield str(root)
    gc.collect()
    shutil.rmtree(root, ignore_errors=True)

def test_tabular_parser_csv(temp_storage):
    parser = TabularParser(storage_root=temp_storage)
    workspace_id = uuid4()
    dataset_id = uuid4()
    version_id = uuid4()
    
    # Create a dummy CSV
    csv_path = os.path.join(temp_storage, "test.csv")
    df = pd.DataFrame({"a": [1, 2], "b": ["x", "y"]})
    df.to_csv(csv_path, index=False)
    
    sheets, profiles = parser.parse_and_materialize(
        csv_path, workspace_id, dataset_id, version_id
    )
    
    assert len(sheets) == 1
    assert sheets[0].name == "default"
    assert len(profiles) == 2
    
    # Test Preview
    preview = parser.get_preview(workspace_id, dataset_id)
    assert preview["rows"] == [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]
    assert preview["columns"] == ["a", "b"]
    
    parser.delete_materialization(workspace_id, dataset_id)

def test_tabular_parser_excel(temp_storage):
    parser = TabularParser(storage_root=temp_storage)
    workspace_id = uuid4()
    dataset_id = uuid4()
    version_id = uuid4()
    
    # Create a dummy Excel with 2 sheets
    xlsx_path = os.path.join(temp_storage, "test.xlsx")
    with pd.ExcelWriter(xlsx_path) as writer:
        pd.DataFrame({"a": [1]}).to_excel(writer, sheet_name="Sheet1", index=False)
        pd.DataFrame({"b": [2]}).to_excel(writer, sheet_name="Sheet2", index=False)
    
    sheets, profiles = parser.parse_and_materialize(
        xlsx_path, workspace_id, dataset_id, version_id
    )
    
    assert len(sheets) == 2
    assert {s.name for s in sheets} == {"Sheet1", "Sheet2"}
    
    # Test Preview Sheet2
    preview = parser.get_preview(workspace_id, dataset_id, sheet_name="Sheet2")
    assert preview["rows"][0]["b"] == 2
    
    parser.delete_materialization(workspace_id, dataset_id)

def test_knowledge_parser_docx_table_only(temp_storage):
    from docx import Document

    parser = KnowledgeParser(gemini_api_key=None)
    docx_path = os.path.join(temp_storage, "table-only.docx")

    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Hạng mục"
    table.cell(0, 1).text = "Kết quả"
    table.cell(1, 0).text = "Thu hồi chứng chỉ"
    table.cell(1, 1).text = "Thành công"
    document.save(docx_path)

    chunk_data = parser.parse_and_chunk(docx_path)

    assert len(chunk_data) > 0
    assert "Thu hồi chứng chỉ" in chunk_data[0]["content"]
    assert "Thành công" in chunk_data[0]["content"]

@pytest.mark.asyncio
async def test_knowledge_parser_mock(temp_storage):
    # Test with no API key (mock/dummy mode)
    parser = KnowledgeParser(gemini_api_key=None)
    txt_path = os.path.join(temp_storage, "test.txt")
    with open(txt_path, "w") as f:
        f.write("Hello world. " * 100) # Enough to trigger potential chunking
        
    workspace_id = uuid4()
    asset_id = uuid4()
    version_id = uuid4()
    
    chunks = await parser.process_file(txt_path, workspace_id, asset_id, version_id)
    assert len(chunks) > 0
    assert chunks[0].content.startswith("Hello world.")
    assert len(chunks[0].embedding) == 768

@pytest.mark.asyncio
async def test_knowledge_parser_gemini_uses_supported_model_and_dimension():
    parser = KnowledgeParser(gemini_api_key="test-key")

    captured: dict[str, object] = {}

    class FakeEmbedding:
        def __init__(self, values):
            self.values = values

    class FakeModels:
        def embed_content(self, *, model, contents, config):
            captured["model"] = model
            captured["contents"] = contents
            captured["config"] = config
            return type("FakeResult", (), {"embeddings": [FakeEmbedding([0.2] * 768)]})()

    parser.client = type("FakeClient", (), {"models": FakeModels()})()

    embeddings = await parser.generate_embeddings([{"content": "hello world", "index": 0}])

    assert captured["model"] == "gemini-embedding-001"
    assert captured["contents"] == ["hello world"]
    assert captured["config"].output_dimensionality == 768
    assert len(embeddings[0]) == 768
